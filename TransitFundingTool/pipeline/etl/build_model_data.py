"""Build the data file for the Regional Transit Funding Model.

Reads the source workbooks named in etl/config.yaml, merges them with the
hand-maintained seed files in etl/seed/, and writes:

  data/model-data.json      the model data
  data/model-data.js        the same data as a script the page can load
  data/build-report.md      what was read, what changed, and any warnings
  dist/transit-funding-model.html   the page with the data inlined (one file)

Workbook values win over the seed, except for fields the seed lists under
`keep`. A kept field stores the workbook value seen when it was kept. If the
workbook later moves away from that value, the report warns so a person can
decide which value is right.

Rows are found by their label text, so inserting rows in a workbook is safe.
Renaming a label the seed points to stops the build with a message naming it.

Run from ArcGIS Pro Python (arcgispro-py3):
    propy etl/build_model_data.py            (or double-click build_data.bat)
    propy etl/build_model_data.py --strict   (fail when there are warnings)
    propy etl/build_model_data.py --check    (build, then compare with the page in model/)
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import hashlib
import json
import math
import re
import sys
from pathlib import Path

try:
    import yaml
except ImportError:  # pragma: no cover
    sys.exit("PyYAML is missing. Run this from ArcGIS Pro Python (propy), which includes it.")
try:
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is missing. Run this from ArcGIS Pro Python (propy), which includes it.")

HERE = Path(__file__).resolve().parent          # pipeline/etl
ROOT = HERE.parent                              # pipeline
SEP = " › "                                # the "›" used in citation strings
YEAR_COLS_DEFAULT = ["G", "H", "I", "J", "K"]


class BuildError(Exception):
    """A problem that stops the build."""


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────
def load_yaml(path: Path):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def num(v):
    """Workbook cell value as a float, or None for blanks and dashes."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("$", "")
    if s in ("", "-", "--", "–", "—"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def text(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def tidy(v):
    """Workbook note text as the page shows it: single spaces, first letter capitalized.
    Blank returns None, which leaves the seed value in place."""
    t = text(v)
    if not t:
        return None
    return t[0].upper() + t[1:]


def same(a, b, tol=1e-6):
    """Equality that tolerates float noise and treats None and missing alike."""
    if isinstance(a, float) and isinstance(b, (int, float)) or isinstance(b, float) and isinstance(a, (int, float)):
        if a is None or b is None:
            return False
        return math.isclose(float(a), float(b), rel_tol=tol, abs_tol=1e-6)
    if isinstance(a, dict) and isinstance(b, dict):
        keys = set(a) | set(b)
        return all(same(a.get(k), b.get(k), tol) for k in keys)
    if isinstance(a, list) and isinstance(b, list):
        return len(a) == len(b) and all(same(x, y, tol) for x, y in zip(a, b))
    return a == b


def cell_value(ws, r, c):
    """Cell value, reading the top-left cell when (r, c) sits inside a merged range."""
    if isinstance(c, str):
        c = column_index_from_string(c)
    v = ws.cell(r, c).value
    if v is None and ws.merged_cells.ranges:
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
                return ws.cell(rng.min_row, rng.min_col).value
    return v


def sha16(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


# ─────────────────────────────────────────────────────────────────────────────
# Workbook access
# ─────────────────────────────────────────────────────────────────────────────
class Workbooks:
    def __init__(self, cfg: dict, base: Path):
        self.cfg = cfg
        folder = Path(cfg["workbook_folder"])
        self.folder = folder if folder.is_absolute() else (base / folder).resolve()
        self._wb = {}
        self.files = {}

    def path(self, key: str) -> Path:
        name = self.cfg["workbooks"][key]["file"]
        p = self.folder / name
        if not p.exists():
            raise BuildError(f"Workbook not found: {p}. Set workbook_folder in etl/config.yaml.")
        return p

    def book(self, key: str):
        if key not in self._wb:
            p = self.path(key)
            self._wb[key] = openpyxl.load_workbook(p, data_only=True)
            self.files[key] = {"file": p.name, "sha256": sha16(p),
                               "modified": dt.datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")}
        return self._wb[key]

    def sheet(self, key: str, name: str):
        wb = self.book(key)
        if name not in wb.sheetnames:
            raise BuildError(f"Sheet '{name}' not found in {self.cfg['workbooks'][key]['file']}.")
        return wb[name]

    def file_name(self, key: str) -> str:
        return self.cfg["workbooks"][key]["file"]

    def row_hits(self, key: str, sheet: str, label: str, col: str | None = None) -> list[int]:
        """Rows, top to bottom, with a cell matching `label` left of `col` (or anywhere when col is None)."""
        ws = self.sheet(key, sheet)
        target = text(label).lower()
        limit = column_index_from_string(col) - 1 if col else ws.max_column
        hits = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max(1, limit)):
            for c in row:
                if c.value is not None and text(c.value).lower() == target:
                    hits.append(c.row)
                    break
        return hits

    def find_row(self, key: str, sheet: str, label: str, col: str | None = None, near: int | None = None,
                 occurrence: int | None = None, under: dict | None = None) -> int:
        """Row whose label cell matches `label`, with the label left of `col`.
        When a label repeats on the sheet, `occurrence` picks the nth match from the top,
        which stays right when rows are inserted. `under` picks the first match below a
        section label. `near` (the row when the seed was written) breaks any remaining tie."""
        hits = self.row_hits(key, sheet, label, col)
        if not hits:
            raise BuildError(f"Row label '{label}' not found on {self.file_name(key)} › {sheet}. "
                             f"If the label was renamed in the workbook, update it in etl/seed.")
        if under:
            top = self.find_row(key, sheet, under["label"], col, under.get("near"), under.get("occurrence"))
            below = [r for r in hits if r > top]
            if not below:
                raise BuildError(f"Row label '{label}' not found below '{under['label']}' on {self.file_name(key)} › {sheet}.")
            return below[0]
        if occurrence:
            if occurrence > len(hits):
                raise BuildError(f"Row label '{label}' appears {len(hits)} time(s) on {self.file_name(key)} › {sheet}; the seed expects occurrence {occurrence}.")
            return hits[occurrence - 1]
        if near:
            hits.sort(key=lambda r: (abs(r - near), r))
        return hits[0]

    def read(self, loc: dict):
        """Read a locator. Returns (value, row, cells) where cells is the A1 text used in citations."""
        key, sheet = loc.get("wb", "plan"), loc["sheet"]
        ws = self.sheet(key, sheet)
        cols = loc.get("cols") or [loc["col"]]
        r = self.find_row(key, sheet, loc["row"], cols[0], loc.get("near"), loc.get("occurrence"), loc.get("under"))
        vals = [ws[f"{c}{r}"].value for c in cols]
        if loc.get("sum"):
            nums = [num(v) for v in vals]
            value = sum(v for v in nums if v is not None) if any(v is not None for v in nums) else None
            cells = f"{cols[0]}{r}:{cols[-1]}{r}"
        else:
            value = vals[0] if len(vals) == 1 else vals
            cells = f"{cols[0]}{r}"
        return value, r, cells


# ─────────────────────────────────────────────────────────────────────────────
# The build
# ─────────────────────────────────────────────────────────────────────────────
class Build:
    def __init__(self, cfg_path: Path):
        self.cfg_path = cfg_path
        self.cfg = load_yaml(cfg_path)
        self.wb = Workbooks(self.cfg, cfg_path.parent)
        seed_dir = cfg_path.parent / self.cfg.get("seed_folder", "seed")
        self.seed = {name: load_yaml(seed_dir / f"{name}.yaml") for name in ("model", "services", "operators", "sources", "fta", "notes")}
        self.warnings: list[str] = []
        self.changes: list[str] = []   # workbook values that differ from the seed's stored value
        self.kept: list[str] = []      # seed values held over the workbook
        self.read_log: list[str] = []

    # a field read from the workbook, with keep handling
    def resolve(self, where: str, field: str, seed_val, wb_val, keep: dict | None):
        k = (keep or {}).get(field)
        if k is not None:
            snap = k.get("workbook") if isinstance(k, dict) else None
            reason = k.get("reason", "") if isinstance(k, dict) else str(k)
            if isinstance(k, dict) and "workbook" in k and not same(snap, wb_val):
                self.warnings.append(f"{where}: `{field}` is kept in the seed ({reason}), but the workbook value moved "
                                     f"from {fmt(snap)} to {fmt(wb_val)}. Review, then update `keep.{field}.workbook` or drop the keep.")
            else:
                self.kept.append(f"{where}: `{field}` = {fmt(seed_val)} (workbook {fmt(wb_val)}). {reason}")
            return seed_val
        if not same(seed_val, wb_val):
            self.changes.append(f"{where}: `{field}` {fmt(seed_val)} → {fmt(wb_val)}")
        return wb_val

    def fname(self, key):
        return self.wb.file_name(key)

    # ── services ──
    def build_services(self):
        cfg = self.cfg["costs_tab"]
        ws = self.wb.sheet(cfg.get("wb", "scen"), cfg["sheet"])
        hdr_row = self.header_row(ws, cfg["label_header"], cfg.get("header_row"))
        headers = {text(ws.cell(hdr_row, c).value).lower(): get_column_letter(c) for c in range(1, ws.max_column + 1) if ws.cell(hdr_row, c).value}

        def col(name):
            c = headers.get(text(name).lower())
            if not c:
                raise BuildError(f"Column '{name}' not found in row {hdr_row} of {self.fname(cfg.get('wb', 'scen'))} › {cfg['sheet']}.")
            return c

        label_c, agency_c = col(cfg["label_header"]), col(cfg["agency_header"])
        fields = cfg["fields"]  # field -> header
        seen_rows = set()
        out = []
        for s in self.seed["services"]["services"]:
            s = copy.deepcopy(s)
            src = s.pop("from", None) or {}
            keep = s.pop("keep", None)
            where = f"Service {s['id']}"
            if "costs" in src:
                lab, ag = src["costs"]["label"], src["costs"]["agency"]
                rows = [r for r in range(hdr_row + 1, ws.max_row + 1)
                        if text(cell_value(ws, r, label_c)).lower() == text(lab).lower() and text(cell_value(ws, r, agency_c)).lower() == text(ag).lower()]
                if not rows:
                    raise BuildError(f"{where}: no row on the costs tab with Service/Route '{lab}' and Agency '{ag}'.")
                r = rows[0]
                seen_rows.add(r)
                for field, header in fields.items():
                    raw = ws[f"{col(header)}{r}"].value
                    if field in ("trips", "vsh", "vsm"):
                        v = num(raw)
                    else:
                        v = tidy(raw)
                        if v is None:          # blank cell: the seed value stands
                            continue
                    s[field] = self.resolve(where, field, s.get(field), v, keep)
                s["cite"] = f"{self.fname(cfg.get('wb', 'scen'))}{SEP}{cfg['sheet']} row {r}"
                self.read_log.append(f"{where}: costs row {r}")
            if "refCost" in src:
                v, r, cells = self.wb.read(src["refCost"])
                s["refCost"] = self.resolve(where, "refCost", s.get("refCost"), round(num(v)) if num(v) is not None else None, keep)
            if "costSchedule" in src:
                loc = dict(src["costSchedule"])
                loc.setdefault("cols", YEAR_COLS_DEFAULT)
                vals, r, _ = self.wb.read(loc)
                sched = [num(v) or 0.0 for v in vals]
                s["costSchedule"] = self.resolve(where, "costSchedule", s.get("costSchedule"), sched, keep)
                s["cite"] = f"{self.fname(loc.get('wb', 'plan'))}{SEP}{loc['sheet']} row {r}"
            out.append(s)
        # rows on the costs tab that no service uses
        ignore = {text(x).lower() for x in cfg.get("ignore_labels", [])}
        for r in range(hdr_row + 1, ws.max_row + 1):
            lab = text(ws[f"{label_c}{r}"].value)
            if lab and r not in seen_rows and lab.lower() not in ignore:
                self.warnings.append(f"Costs tab row {r} ('{lab}', {text(ws[f'{agency_c}{r}'].value)}) is not in etl/seed/services.yaml. "
                                     f"Add it as a service, or list it under costs_tab.ignore_labels in config.yaml.")
        return out

    def header_row(self, ws, label, hint=None):
        """Row that holds the column headers, found by a header label so rows inserted above it are safe."""
        target = text(label).lower()
        rows = [r for r in range(1, min(ws.max_row, 60) + 1)
                if any(text(ws.cell(r, c).value).lower() == target for c in range(1, min(ws.max_column, 60) + 1))]
        if not rows:
            raise BuildError(f"Header '{label}' not found in the first 60 rows of {ws.title}.")
        return min(rows, key=lambda r: abs(r - (hint or r)))

    # ── operators ──
    def build_operators(self):
        out = []
        for o in self.seed["operators"]["operators"]:
            o = copy.deepcopy(o)
            src = o.pop("from", None)
            keep = o.pop("keep", None)
            if src:
                v, r, cells = self.wb.read(src)
                o["rate"] = self.resolve(f"Operator {o['id']}", "rate", o.get("rate"), num(v), keep)
                o["cite"] = f"{self.fname(src.get('wb', 'scen'))}{SEP}{src['sheet']} {cells}{src.get('cite_suffix', '')}"
            out.append(o)
        return out

    # ── Revenue %s ──
    def revenue_row(self, label: str, near: int | None, occurrence: int | None = None):
        cfg = self.cfg["revenue_tab"]
        key, sheet = cfg.get("wb", "plan"), cfg["sheet"]
        r = self.wb.find_row(key, sheet, label, get_column_letter(column_index_from_string(cfg["label_col"]) + 1), near, occurrence)
        ws = self.wb.sheet(key, sheet)
        cond = text(ws[f"{cfg['condition_col']}{r}"].value)
        match = text(ws[f"{cfg['match_col']}{r}"].value)
        shares = {}
        unmapped = []
        hdr = self.header_row(ws, cfg["header_label"], cfg.get("header_row"))
        any_share = False
        for c in range(column_index_from_string(cfg["first_route_col"]), column_index_from_string(cfg["last_route_col"]) + 1):
            h = text(ws.cell(hdr, c).value)
            v = num(ws.cell(r, c).value)
            if v is None:
                continue
            any_share = any_share or v != 0
            sid = cfg["columns"].get(h, "__missing__")
            if sid == "__missing__":
                if v:
                    unmapped.append(h)
                continue
            if sid is None:          # mapped to nothing on purpose in config.yaml
                continue
            for one in (sid if isinstance(sid, list) else [sid]):
                shares[one] = shares.get(one, 0.0) + v
        return {"row": r, "cond": cond, "match": match, "shares": shares if any_share else None, "unmapped": unmapped}

    # ── funding sources ──
    def build_lines(self, services):
        svc_agency = {s["id"]: s["agency"] for s in services}
        rev_cfg = self.cfg["revenue_tab"]
        rev_file, rev_sheet = self.fname(rev_cfg.get("wb", "plan")), rev_cfg["sheet"]
        amounts = {}
        out = []
        pending_minus = []
        for l in self.seed["sources"]["sources"]:
            l = copy.deepcopy(l)
            src = l.pop("from", None) or {}
            keep = l.pop("keep", None)
            where = f"Source {l['id']}"
            parts = []
            if "amount" in src:
                loc = src["amount"]
                v, r, cells = self.wb.read(loc)
                val = num(v) if not isinstance(v, list) else None
                if loc.get("sum"):
                    val = v
                if loc.get("minus"):
                    pending_minus.append((l, loc, val, keep, where))
                else:
                    l["amount"] = self.resolve(where, "amount", l.get("amount"), val if val is not None else 0.0, keep)
                p = f"FY 26-27 amount: {self.fname(loc.get('wb', 'plan'))}{SEP}{loc['sheet']} {cells}"
                if loc.get("minus"):
                    p += "; minus amount of " + ", ".join(loc["minus"])
                parts.append(p)
            if "years" in src:
                loc = src["years"]
                cols = loc["cols"]
                vals, r, _ = self.wb.read(loc)
                vals = vals if isinstance(vals, list) else [vals]
                ov = list(l.get("overrides") or [None] * 5)
                new = list(ov)
                for c, v in zip(cols, vals):
                    idx = YEAR_COLS_DEFAULT.index(c)
                    new[idx] = num(v)
                l["overrides"] = self.resolve(where, "overrides", ov, new, keep)
                parts.append(f"Year values: {self.fname(loc.get('wb', 'plan'))}{SEP}{loc['sheet']} row {r} ({', '.join(cols)})")
            if "revenue" in src:
                rv = src["revenue"]
                info = self.revenue_row(rv["row"], rv.get("near"), rv.get("occurrence"))
                r = info["row"]
                pct = {sid: round(v, 12) for sid, v in (info["shares"] or {}).items() if v and svc_agency.get(sid) == agency_of(l)}
                # a row with no shares for this agency's services leaves the seed allocation in place
                k = (keep or {}).get("alloc")
                if rv.get("shares", True) and (pct or k is not None):
                    if info["unmapped"] and k is None:
                        self.warnings.append(f"{where}: Revenue %s row {r} has shares in columns with no service mapping ({', '.join(info['unmapped'])}). "
                                             f"Map them under revenue_tab.columns in config.yaml.")
                    if k is not None:
                        snap = k.get("workbook") if isinstance(k, dict) else None
                        if isinstance(k, dict) and "workbook" in k and not same(snap, pct):
                            self.warnings.append(f"{where}: shares are kept in the seed ({k.get('reason', '')}), but Revenue %s row {r} changed. "
                                                 f"Review, then update `keep.alloc.workbook` or drop the keep.")
                        else:
                            self.kept.append(f"{where}: shares held in the seed. {k.get('reason', '') if isinstance(k, dict) else k}")
                        parts.append(f"Shares: seed allocation (Revenue %s row {r} overridden; see build report)")
                    else:
                        old = (l.get("alloc") or {}).get("pct") or {}
                        if not same(old, pct, 1e-9):
                            self.changes.append(f"{where}: shares from Revenue %s row {r} changed")
                        l["alloc"] = {"mode": "manual", "pct": pct}
                        parts.append(f"Shares: {rev_file}{SEP}{rev_sheet} row {r}")
                if rv.get("match", True):
                    for field, raw in (("conditionText", info["cond"]), ("matchText", info["match"])):
                        v = tidy(raw)
                        if v is not None:      # blank cell: the seed value stands
                            l[field] = self.resolve(where, field, l.get(field), v, keep)
                    if l.get("level") == "federal":
                        fs = fed_share_from(info["match"])
                        if fs is not None:
                            l["fedShare"] = self.resolve(where, "fedShare", l.get("fedShare"), fs, keep)
                    parts.append(f"Conditions and match: {rev_file}{SEP}{rev_sheet} row {r}, columns {rev_cfg['condition_col']} and {rev_cfg['match_col']}")
            if parts:
                l.setdefault("cites", {})["build"] = " | ".join(parts)
            amounts[l["id"]] = l
            out.append(l)
        for l, loc, val, keep, where in pending_minus:
            sub = sum(amounts[m].get("amount") or 0 for m in loc["minus"])
            l["amount"] = self.resolve(where, "amount", l.get("amount"), (val or 0.0) - sub, keep)
        # Revenue %s rows nobody reads
        self.check_revenue_coverage()
        return out

    def check_revenue_coverage(self):
        cfg = self.cfg["revenue_tab"]
        key, sheet = cfg.get("wb", "plan"), cfg["sheet"]
        ws = self.wb.sheet(key, sheet)
        used = set()
        for l in self.seed["sources"]["sources"]:
            rv = ((l.get("from") or {}).get("revenue") or {})
            if rv:
                used.add(text(rv["row"]).lower())
        ignore = {text(x).lower() for x in cfg.get("ignore_rows", [])}
        hdr = self.header_row(ws, cfg["header_label"], cfg.get("header_row"))
        stop = ws.max_row
        if cfg.get("end_label"):
            stop = self.wb.find_row(key, sheet, cfg["end_label"], get_column_letter(column_index_from_string(cfg["label_col"]) + 1)) - 1
        for r in range(hdr + 1, stop + 1):
            lab = text(ws[f"{cfg['label_col']}{r}"].value)
            if lab and lab.lower() not in used and lab.lower() not in ignore:
                self.warnings.append(f"Revenue %s row {r} ('{lab}') is not read by any source. Add a source with `from.revenue` "
                                     f"in etl/seed/sources.yaml, or list it under revenue_tab.ignore_rows in config.yaml.")
        for c in range(column_index_from_string(cfg["first_route_col"]), column_index_from_string(cfg["last_route_col"]) + 1):
            h = text(ws.cell(hdr, c).value)
            if h and h not in cfg["columns"]:
                self.warnings.append(f"Revenue %s column {get_column_letter(c)} ('{h}') has no service mapping. Add it under revenue_tab.columns in config.yaml.")

    # ── FTA sub-allocation ──
    def build_fta(self):
        f = copy.deepcopy(self.seed["fta"]["fta"])
        locs = self.seed["fta"].get("from", {})
        keep = self.seed["fta"].get("keep")
        for path, loc in locs.items():
            v, r, cells = self.wb.read(loc)
            cur = get_path(f, path)
            set_path(f, path, self.resolve("FTA split", path, cur, num(v), keep))
        return f

    # ── assemble ──
    def run(self, build_id: str | None = None):
        m = self.seed["model"]
        services = self.build_services()
        operators = self.build_operators()
        lines = self.build_lines(services)
        fta = self.build_fta()
        self.validate(services, operators, lines)
        now = dt.datetime.now()
        bid = build_id or now.strftime("%Y%m%d") + "a"
        data = {
            "schema": m["schema"], "modelVersion": m["modelVersion"],
            "sourceFiles": m["sourceFiles"], "sourceFileNotes": m["sourceFileNotes"],
            "agencies": m["agencies"], "jurisdictions": m["jurisdictions"],
            "services": services, "operators": operators, "lines": lines, "fta": fta,
            "dataNotes": self.seed["notes"]["dataNotes"], "agreements": self.seed["notes"]["agreements"],
            "build": {"id": bid, "builtAt": now.isoformat(timespec="seconds"),
                      "workbooks": {k: v for k, v in self.wb.files.items()}, "warnings": list(self.warnings)},
        }
        return data

    def validate(self, services, operators, lines):
        sids = {s["id"] for s in services}
        ops = {o["id"] for o in operators}
        ags = {a["id"] for a in self.seed["model"]["agencies"]}
        jur = {j["id"] for j in self.seed["model"]["jurisdictions"]}
        for s in services:
            if s["operator"] not in ops:
                self.warnings.append(f"Service {s['id']}: operator '{s['operator']}' is not in operators.yaml.")
            if s["agency"] not in ags:
                self.warnings.append(f"Service {s['id']}: agency '{s['agency']}' is not in model.yaml.")
            for j in (s.get("juris") or {}):
                if j not in jur:
                    self.warnings.append(f"Service {s['id']}: jurisdiction '{j}' is not in model.yaml.")
            t = sum((s.get("juris") or {}).values())
            if s.get("include") and not s.get("plan") and abs(t - 1) > 0.005:
                self.warnings.append(f"Service {s['id']}: jurisdiction shares total {t:.3f}, not 1.")
        for l in lines:
            if l["recipient"] not in ags:
                self.warnings.append(f"Source {l['id']}: receiving agency '{l['recipient']}' is not in model.yaml.")
            pct = (l.get("alloc") or {}).get("pct") or {}
            for sid in pct:
                if sid not in sids:
                    self.warnings.append(f"Source {l['id']}: share for unknown service '{sid}'.")
            if sum(pct.values()) > 1.0005:
                self.warnings.append(f"Source {l['id']}: shares total {sum(pct.values()):.4f}, more than 100 percent.")
            r = l.get("rules") or {}
            for sid in (r.get("routes") or []) + (r.get("preferRoutes") or []):
                if sid not in sids:
                    self.warnings.append(f"Source {l['id']}: rule names unknown service '{sid}'.")


def agency_of(line):
    return line["recipient"]


def fed_share_from(match: str):
    """'50% non-federal' → 0.5 federal share; 'no match requirement' → None (no change)."""
    m = re.search(r"([\d.]+)\s*%\s*non-?federal", match or "", re.I)
    if not m:
        return None
    return round(1 - float(m.group(1)) / 100, 6)


def get_path(o, path):
    for k in path.split("."):
        o = None if o is None else o.get(k)
    return o


def set_path(o, path, v):
    ks = path.split(".")
    for k in ks[:-1]:
        o = o.setdefault(k, {})
    o[ks[-1]] = v


def fmt(v):
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if abs(v) < 1e6 or v != int(v) else f"{v:,.0f}"
    if isinstance(v, (dict, list)):
        s = json.dumps(v)
        return s if len(s) < 90 else s[:87] + "..."
    if v is None:
        return "blank"
    return repr(v)


# ─────────────────────────────────────────────────────────────────────────────
# Output
# ─────────────────────────────────────────────────────────────────────────────
def write_outputs(build: Build, data: dict):
    out_cfg = build.cfg["output"]
    base = build.cfg_path.parent
    data_dir = (base / out_cfg["data_folder"]).resolve()
    data_dir.mkdir(parents=True, exist_ok=True)
    js = json.dumps(data, indent=1, ensure_ascii=False)
    (data_dir / "model-data.json").write_text(js, encoding="utf-8")
    (data_dir / "model-data.js").write_text("window.TRANSIT_MODEL_DATA = " + js + ";\n", encoding="utf-8")
    # report
    rep = [f"# Data build {data['build']['id']}", "", f"Built {data['build']['builtAt']}.", "", "## Workbooks", ""]
    for k, w in data["build"]["workbooks"].items():
        rep.append(f"- {w['file']}, saved {w['modified']}, sha256 {w['sha256']}")
    rep += ["", f"## Warnings ({len(build.warnings)})", ""] + ([f"- {w}" for w in build.warnings] or ["None."])
    rep += ["", f"## Workbook values that differ from the seed ({len(build.changes)})", "",
            "These came from the workbooks and replace the values stored in etl/seed. Copy them into the seed when you agree with them, so the seed stays a true record.", ""]
    rep += [f"- {c}" for c in build.changes] or ["None."]
    rep += ["", f"## Seed values held over the workbook ({len(build.kept)})", ""] + ([f"- {c}" for c in build.kept] or ["None."])
    (data_dir / "build-report.md").write_text("\n".join(rep) + "\n", encoding="utf-8")
    # single-file page
    tpl = (base / out_cfg["template"]).resolve()
    dist = (base / out_cfg["dist"]).resolve()
    html = tpl.read_text(encoding="utf-8")
    tag = '<script src="data/model-data.js"></script>'
    if tag not in html:
        raise BuildError(f"The template must contain {tag} where the data goes.")
    html = html.replace(tag, "<script>\nwindow.TRANSIT_MODEL_DATA = " + js + ";\n</script>")
    dist.parent.mkdir(parents=True, exist_ok=True)
    dist.write_text(html, encoding="utf-8")
    return data_dir, dist


def page_data(path: Path) -> dict:
    s = path.read_text(encoding="utf-8")
    i = s.index("window.TRANSIT_MODEL_DATA")
    j = s.index("</script>", i)
    body = s[i:j]
    return json.loads(body[body.index("{"):body.rindex("}") + 1])


def compare(a, b, path="", out=None, tol=1e-6):
    out = [] if out is None else out
    if isinstance(a, dict) and isinstance(b, dict):
        for k in sorted(set(a) | set(b)):
            compare(a.get(k), b.get(k), f"{path}.{k}" if path else k, out, tol)
    elif isinstance(a, list) and isinstance(b, list):
        if len(a) != len(b):
            out.append(f"{path}: length {len(a)} vs {len(b)}")
        for i, (x, y) in enumerate(zip(a, b)):
            ident = x.get("id") if isinstance(x, dict) and "id" in x else i
            compare(x, y, f"{path}[{ident}]", out, tol)
    elif not same(a, b, tol):
        out.append(f"{path}: {fmt(a)} vs {fmt(b)}")
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default=str(HERE / "config.yaml"))
    ap.add_argument("--strict", action="store_true", help="exit with an error when there are warnings")
    ap.add_argument("--check", metavar="PAGE", nargs="?", const="", help="compare the built data with a page (default: the page named in config)")
    ap.add_argument("--build-id", help="set the build id instead of today's date")
    args = ap.parse_args(argv)
    try:
        b = Build(Path(args.config).resolve())
        data = b.run(args.build_id)
        data_dir, dist = write_outputs(b, data)
    except BuildError as e:
        print(f"BUILD FAILED: {e}", file=sys.stderr)
        return 2
    print(f"Build {data['build']['id']}: {len(data['services'])} services, {len(data['lines'])} sources.")
    print(f"  wrote {data_dir / 'model-data.json'}")
    print(f"  wrote {dist}")
    print(f"  {len(b.warnings)} warning(s), {len(b.changes)} workbook change(s), {len(b.kept)} kept value(s). See {data_dir / 'build-report.md'}.")
    for w in b.warnings:
        print("  WARNING:", w)
    if args.check is not None:
        page = Path(args.check) if args.check else (b.cfg_path.parent / b.cfg["output"]["check_page"]).resolve()
        diffs = compare({k: v for k, v in data.items() if k != "build"}, {k: v for k, v in page_data(page).items() if k != "build"})
        print(f"Compared with {page}: {len(diffs)} difference(s).")
        for d in diffs[:200]:
            print("  ", d)
        if diffs:
            return 3
    if args.strict and b.warnings:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
