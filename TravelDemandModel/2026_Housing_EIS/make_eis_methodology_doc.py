"""
Produce a cleaned-up, Appendix-F-styled version of
'Notes for TDM documentation.docx' for the 2026 Housing EIS.

Style reference values measured from Appendix F- Data and Forecasting.docx:
  TRPA Green  #54746A  – H1/H3/H4 text; H2 background
  White       #FFFFFF  – H2 text
  Dark Navy   #222B35  – table header fill
  Caption     #44546A  – caption text color
  Row Gray    #A6A6A6  – alternating table row fill
  Font        Calibri, 11 pt body
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = Path(r"C:\Users\amcclary\Documents\GitHub\Transportation\TravelDemandModel\2026_Housing_EIS\EIS_Data_Forecasting_Methodology.docx")

# ── palette ───────────────────────────────────────────────────────────────────
TRPA_GREEN  = "54746A"
WHITE       = "FFFFFF"
DARK_NAVY   = "222B35"
CAPTION_CLR = "44546A"
ROW_GRAY    = "A6A6A6"
NOTE_GRAY   = "595959"
TOTAL_BLUE  = "D9E1F2"

# ── low-level helpers ─────────────────────────────────────────────────────────
def _shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_fill)
    tcPr.append(shd)

def _shade_para(p, hex_fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_fill)
    pPr.append(shd)

def _col_width(cell, inches):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW  = OxmlElement("w:tcW")
    tcW.set(qn("w:w"),    str(int(inches * 1440)))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)

def _fill(cell, text, size=10, bold=False, color=None, align=None, bg=None, width=None):
    cell.text = str(text)
    p   = cell.paragraphs[0]
    run = p.runs[0]
    run.font.size = Pt(size)
    run.bold      = bold
    if color:  run.font.color.rgb = RGBColor.from_string(color)
    if align:  p.alignment        = align
    if bg:     _shade_cell(cell, bg)
    if width:  _col_width(cell, width)

# ── heading / body helpers ────────────────────────────────────────────────────
def h1(doc, text):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold            = True
    run.font.size       = Pt(20)
    run.font.color.rgb  = RGBColor.from_string(TRPA_GREEN)
    run.font.name       = "Calibri"
    pf = p.paragraph_format
    pf.space_before = Pt(18)
    pf.space_after  = Pt(6)
    return p

def h2(doc, text):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold            = True
    run.font.size       = Pt(14)
    run.font.color.rgb  = RGBColor.from_string(WHITE)
    run.font.name       = "Calibri"
    _shade_para(p, TRPA_GREEN)
    pf = p.paragraph_format
    pf.space_before = Pt(14)
    pf.space_after  = Pt(6)
    # left/right indent so the background fill fills the full width
    pf.left_indent  = Inches(-1)
    pf.right_indent = Inches(-1)
    # but keep text indented normally
    pf.first_line_indent = Inches(1)
    return p

def h2_simple(doc, text):
    """H2 without margin bleed — simpler, same color scheme."""
    p   = doc.add_paragraph()
    run = p.add_run(f"  {text}")
    run.bold            = True
    run.font.size       = Pt(14)
    run.font.color.rgb  = RGBColor.from_string(WHITE)
    run.font.name       = "Calibri"
    _shade_para(p, TRPA_GREEN)
    pf = p.paragraph_format
    pf.space_before = Pt(14)
    pf.space_after  = Pt(6)
    return p

def h3(doc, text):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold            = True
    run.font.size       = Pt(13)
    run.font.color.rgb  = RGBColor.from_string(TRPA_GREEN)
    run.font.name       = "Calibri"
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after  = Pt(4)
    return p

def h4(doc, text):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold            = True
    run.font.size       = Pt(11)
    run.font.color.rgb  = RGBColor.from_string(TRPA_GREEN)
    run.font.name       = "Calibri"
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after  = Pt(2)
    return p

def body(doc, text, space_after=6):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    p.paragraph_format.space_after = Pt(space_after)
    return p

def caption(doc, text):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold            = True
    run.font.size       = Pt(9)
    run.font.color.rgb  = RGBColor.from_string(CAPTION_CLR)
    run.font.name       = "Calibri"
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after  = Pt(3)
    return p

def footnote(doc, text):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size      = Pt(8.5)
    run.font.name      = "Calibri"
    run.font.color.rgb = RGBColor.from_string(NOTE_GRAY)
    p.paragraph_format.space_before = Pt(4)
    return p

# ── table builder ─────────────────────────────────────────────────────────────
def build_table(doc, headers, widths, rows, center_data_cols=None, footnote_text=None):
    """
    headers       : list of header strings
    widths        : list of column widths in inches
    rows          : list of (values…) tuples for data rows
    center_data_cols : column indices to center-align in data rows
    """
    center_data_cols = set(center_data_cols or [])
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style     = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # header row
    for i, (cell, hdr) in enumerate(zip(tbl.rows[0].cells, headers)):
        _fill(cell, hdr,
              size=10, bold=True, color=WHITE,
              align=WD_ALIGN_PARAGRAPH.CENTER,
              bg=DARK_NAVY, width=widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # data rows
    for ri, row_vals in enumerate(rows):
        is_total = isinstance(row_vals, dict) and row_vals.get("_total")
        if is_total:
            row_vals = row_vals["values"]
        bg = TOTAL_BLUE if is_total else (ROW_GRAY if ri % 2 == 0 else None)
        new_row = tbl.add_row()
        for j, (cell, val) in enumerate(zip(new_row.cells, row_vals)):
            align = WD_ALIGN_PARAGRAPH.CENTER if j in center_data_cols else None
            bold  = is_total
            _fill(cell, val, size=10, bold=bold, align=align, bg=bg, width=widths[j])

    if footnote_text:
        doc.add_paragraph()
        footnote(doc, footnote_text)

    return tbl


# ══════════════════════════════════════════════════════════════════════════════
# Document
# ══════════════════════════════════════════════════════════════════════════════
doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

s = doc.sections[0]
s.left_margin   = Inches(1)
s.right_margin  = Inches(1)
s.top_margin    = Inches(1)
s.bottom_margin = Inches(1)

# ── Title ─────────────────────────────────────────────────────────────────────
h1(doc, "Appendix F: Data and Forecasting Methodology")

# ══════════════════════════════════════════════════════════════════════════════
# 1. Introduction
# ══════════════════════════════════════════════════════════════════════════════
h2_simple(doc, "Introduction")

body(doc,
    "In order to investigate the impact of proposed scenarios for the Cultivating Community "
    "effort to increase workforce housing, three different travel demand model (TDM) scenarios "
    "were developed. These scenarios used most of the same base assumptions as the 2025 Regional "
    "Transportation Plan (RTP), with the primary differences being updated information about "
    "development that has taken place since the 2022 base year used for that RTP, and the "
    "additional housing and occupancy policies proposed under Alternatives 2 and 3, which are "
    "described in detail below. For a detailed description of the assumptions used for the 2025 "
    "RTP, please refer to the original memo. Both employment and school enrollment were assumed "
    "to increase slightly in Alternatives 2 and 3 in proportion to the increases in population."
)

# ══════════════════════════════════════════════════════════════════════════════
# 2. Development Scenario Descriptions
# ══════════════════════════════════════════════════════════════════════════════
h2_simple(doc, "Development Scenario Descriptions")

# ── Alt 1 ─────────────────────────────────────────────────────────────────────
h3(doc, "Alternative 1 — No Action (Base Year RTP Forecast)")

body(doc,
    "Alternative 1 assumed no change in existing policies and used the same growth assumptions "
    "outlined in the 2025 RTP. Minor differences from the 2022 TDM results were due to starting "
    "from a 2025 base year, which included additional information about the location of residential "
    "units in the Tahoe Basin and development that occurred between 2022 and 2025. Vacation home "
    "rental (VHR) numbers were also adjusted to reflect updated policies in effect at the time "
    "the model was run."
)

# ── Forecast Methodology ──────────────────────────────────────────────────────
h3(doc, "Forecast Methodology")

body(doc,
    "For each alternative, new residential development was allocated spatially to parcels meeting "
    "eligibility criteria specific to each unit pool category. Parcels were selected based on "
    "attributes including existing land use, ownership type, IPES score, housing zoning "
    "designation, location relative to TRPA boundaries and Town Centers, and parcel size. "
    "Unit allocation within each pool proceeded sequentially until the pool target was reached. "
    "Occupancy rates and household income distributions were then applied by pool category to "
    "translate new residential units into occupied households and population estimates for input "
    "to the travel demand model."
)

# ── Alt 2 ─────────────────────────────────────────────────────────────────────
h3(doc, "Alternative 2 — Affordable Housing Focus")

body(doc,
    "Alternative 2 proposes a variety of changes to housing policy designed to increase the "
    "number of residential units constructed in the Tahoe Basin. Under this alternative, "
    "additional housing would be constructed across several unit pool categories, each with "
    "different assumptions about where units would be located. For the purposes of modeling, "
    "new development was assigned to parcels meeting the eligibility criteria for each pool."
)

body(doc,
    "Additional bonus units in this alternative were distributed across the following "
    "categories and location eligibility criteria:",
    space_after=4
)

# ── Bonus Units Table ─────────────────────────────────────────────────────────
bonus_headers = ["Unit Category", "New Units", "Location / Eligibility Criteria"]
bonus_widths  = [1.55, 0.85, 4.1]
bonus_rows    = [
    ("Affordable",           "456", "Within designated Town Centers"),
    ("Moderate",             "285", "Within the Bonus Unit Boundary"),
    ("Achievable",           "228", "Within the Bonus Unit Boundary"),
    ("Achievable",           "57",  "Anywhere residential development is allowed"),
    ("Affordable by Design", "114", "Anywhere residential development is allowed"),
    ("JADU",                 "98",  "Anywhere residential development is allowed"),
    {"_total": True, "values": ("Total Additional Bonus Units", "1,238", "")},
]

build_table(doc, bonus_headers, bonus_widths, bonus_rows,
            center_data_cols=[1],
            footnote_text=(
                "Note: These unit categories are in addition to the Bonus, General, and ADU "
                "pools shared across all alternatives. See Table 1 for full unit counts."
            ))

doc.add_paragraph()  # spacer

# ── Alt 3 ─────────────────────────────────────────────────────────────────────
h3(doc, "Alternative 3 — Housing Activation")

body(doc,
    "Alternative 3 proposes a variety of changes to housing policy intended to increase the "
    "proportion of existing housing in the basin occupied by year-round residents. Under this "
    "alternative, new housing development was assigned to parcels in the same manner as "
    "Alternative 1 (No Action), and the residential occupancy rate was increased uniformly "
    "throughout the basin so that the regional population reached the targeted number of "
    "additional full-time residents."
)

# ══════════════════════════════════════════════════════════════════════════════
# 3. Alternative Comparison Table
# ══════════════════════════════════════════════════════════════════════════════
h2_simple(doc, "Alternative Descriptions and Key Parameters")

caption(doc, "Table 1: Alternative Descriptions and Key Parameters")

alt_headers = [
    "Parameter",
    "Alternative 1\n(Baseline)",
    "Alternative 2\n(Affordable\nHousing Focus)",
    "Alternative 3\n(Housing\nActivation)",
]
alt_widths = [2.6, 1.3, 1.35, 1.35]

alt_rows = [
    ("Total New Residential Units",      "3,377",  "4,615",  "3,377"),
    ("  Bonus Units",                    "1,246",  "1,246",  "1,246"),
    ("  General Units",                  "2,021",  "2,021",  "2,021"),
    ("  ADU",                            "110",    "110",    "110"),
    ("  Affordable Units",               "—",      "456",    "—"),
    ("  Moderate Units",                 "—",      "285",    "—"),
    ("  Achievable Bonus",               "—",      "228",    "—"),
    ("  Achievable General",             "—",      "57",     "—"),
    ("  Affordable by Design",           "—",      "114",    "—"),
    ("  JADU",                           "—",      "98",     "—"),
    ("Additional Units vs. Alternative 1","—",     "+1,238", "—"),
    ("Occupancy Adjustment (2035)",      "None",   "None",   "+332 occupied units"),
    ("Occupancy Adjustment (2050)",      "None",   "None",   "+855 occupied units"),
    ("Town Center Employment Add-back",  "No",     "Yes",    "Yes"),
    ("Population Target (2035)",         "55,592", "Modeled","Modeled"),
    ("Population Target (2050)",         "57,611", "Modeled","Modeled"),
]

build_table(doc, alt_headers, alt_widths, alt_rows,
            center_data_cols=[1, 2, 3],
            footnote_text=(
                "Note: Population targets for Alternative 1 reflect 2025 RTP projections. "
                "Alternatives 2 and 3 model population as an output based on new unit construction "
                "and occupancy assumptions. Dashes (—) indicate no units of that type in the alternative."
            ))

doc.save(OUTPUT)
print(f"Saved: {OUTPUT}")
